"""Build the Gabor theory report and reproducible illustrations for this repository."""
from __future__ import annotations

import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]

import cv2
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
# Load the compatible installed Pillow before the test dependency folder,
# which also contains Pillow binaries for a different Python version.
from PIL import Image
sys.path.insert(0, str(ROOT / ".test-deps"))
sys.path.insert(0, str(ROOT))
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from app.search.config import DEFAULT_CONFIG as CFG

OUT = ROOT / "ouput" / "documents" / "Bao_cao_ly_thuyet_Gabor.docx"
ASSETS = ROOT / "tmp" / "gabor_theory"
ASSETS.mkdir(parents=True, exist_ok=True)
OUT.parent.mkdir(parents=True, exist_ok=True)


def kernel(wavelength, orientation):
    k = cv2.getGaborKernel((CFG.gabor_kernel_size,) * 2,
        CFG.gabor_sigma_ratio * wavelength, orientation * np.pi / CFG.gabor_orientations,
        wavelength, CFG.gabor_gamma, 0, ktype=cv2.CV_32F)
    k -= float(k.mean())
    k /= float(np.linalg.norm(k)) + CFG.l2_eps
    return k


def figures():
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10})
    fig, axes = plt.subplots(4, 4, figsize=(7.2, 6.0))
    for s, lam in enumerate(CFG.gabor_wavelengths):
        for t in range(4):
            k = kernel(lam, t)
            v = np.abs(k).max()
            axes[s, t].imshow(k, cmap="RdBu_r", vmin=-v, vmax=v,
                              extent=(-10.5, 10.5, 10.5, -10.5))
            axes[s, t].set_xticks([])
            axes[s, t].set_yticks([])
            if s == 0:
                axes[s, t].set_title(f"θ = {t * 45}°", pad=8)
            if t == 0:
                axes[s, t].set_ylabel(f"λ = {lam:g} px", labelpad=8)
    fig.tight_layout(pad=0.6)
    fig.savefig(ASSETS / "kernel_bank.png", dpi=220, bbox_inches="tight")
    plt.close(fig)

    y, x = np.mgrid[0:128, 0:128]
    texture = (0.5 + 0.5 * np.cos(2 * np.pi * x / 8)).astype(np.float32)
    fig, axes = plt.subplots(1, 3, figsize=(7.2, 2.2))
    images = [texture, np.abs(cv2.filter2D(texture, cv2.CV_32F, kernel(8, 0))),
              np.abs(cv2.filter2D(texture, cv2.CV_32F, kernel(8, 2)))]
    titles = ["Vân dọc có chu kỳ 8 px", "Đáp ứng θ = 0°", "Đáp ứng θ = 90°"]
    for ax, a, title in zip(axes, images, titles):
        ax.imshow(a, cmap="gray", vmin=0, vmax=max(images[1].max(), 1))
        ax.set_title(title, fontsize=12)
        ax.axis("off")
    fig.tight_layout()
    fig.savefig(ASSETS / "orientation_demo.png", dpi=220, bbox_inches="tight")
    plt.close(fig)


doc = Document()
# Remove decorative defaults and theme font overrides from the base template.
for style in doc.styles:
    if style._element.pPr is not None:
        for border in list(style._element.pPr.findall(qn("w:pBdr"))):
            style._element.pPr.remove(border)
    if style._element.rPr is not None:
        fonts = style._element.rPr.find(qn("w:rFonts"))
        if fonts is not None:
            for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                fonts.attrib.pop(qn("w:" + attr), None)
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
sec.left_margin, sec.right_margin = Cm(2.3), Cm(2.0)
sec.footer_distance = Cm(0.8)
for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"):
    st = doc.styles[name]
    st.font.name = "Times New Roman"
    st.font.color.rgb = RGBColor(0, 0, 0)
    st._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal = doc.styles["Normal"]
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.widow_control = True
for name, size in (("Title", 23), ("Heading 1", 16), ("Heading 2", 13)):
    st = doc.styles[name]
    st.font.size = Pt(size)
    st.font.bold = True
    st.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
    st.paragraph_format.space_after = Pt(8)
    st.paragraph_format.keep_with_next = True
doc.styles["Caption"].font.size = Pt(10.5)
doc.styles["Caption"].font.italic = True
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
footer._p.append(fld)


def p(text, style=None):
    para = doc.add_paragraph(text, style=style)
    if style is None:
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return para


def h(text, level=1):
    return doc.add_heading(text, level)


def page(title):
    doc.add_page_break()
    h(title)


def mr(text):
    r = OxmlElement("m:r")
    rp = OxmlElement("m:rPr")
    sty = OxmlElement("m:sty")
    sty.set(qn("m:val"), "p")
    rp.append(sty)
    r.append(rp)
    t = OxmlElement("m:t")
    t.text = text
    r.append(t)
    return r


def nodes(target, items):
    for item in items if isinstance(items, (list, tuple)) else [items]:
        target.append(mr(item) if isinstance(item, str) else item)


def frac(a, b):
    f = OxmlElement("m:f")
    for tag, value in (("m:num", a), ("m:den", b)):
        part = OxmlElement(tag)
        nodes(part, value)
        f.append(part)
    return f


def sup(base, power):
    e = OxmlElement("m:sSup")
    for tag, value in (("m:e", base), ("m:sup", power)):
        part = OxmlElement(tag)
        nodes(part, value)
        e.append(part)
    return e


def sub(base, index):
    e = OxmlElement("m:sSub")
    for tag, value in (("m:e", base), ("m:sub", index)):
        part = OxmlElement(tag)
        nodes(part, value)
        e.append(part)
    return e


def root(value):
    rad = OxmlElement("m:rad")
    pr = OxmlElement("m:radPr")
    hide = OxmlElement("m:degHide")
    hide.set(qn("m:val"), "1")
    pr.append(hide)
    rad.append(pr)
    rad.append(OxmlElement("m:deg"))
    e = OxmlElement("m:e")
    nodes(e, value)
    rad.append(e)
    return rad


def eq(items, number):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(5)
    para.paragraph_format.space_after = Pt(9)
    math = OxmlElement("m:oMath")
    nodes(math, items)
    para._p.append(math)
    para.add_run(f"     ({number})").font.size = Pt(10)
    return para


def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, w in zip(t.columns, widths):
        c.width = Cm(w)
    for i, text in enumerate(headers):
        t.rows[0].cells[i].text = text
    for row in rows:
        for c, text in zip(t.add_row().cells, row):
            c.text = str(text)
    for ri, row in enumerate(t.rows):
        pr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        pr.append(cant)
        if ri == 0:
            pr.append(OxmlElement("w:tblHeader"))
        for ci, cell in enumerate(row.cells):
            cell.width = Cm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                b = OxmlElement("w:" + edge)
                for attr, value in (("val", "single"), ("sz", "5"), ("color", "D9D9D9")):
                    b.set(qn("w:" + attr), value)
                borders.append(b)
            tcpr.append(borders)
            mar = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                v = OxmlElement("w:" + edge)
                v.set(qn("w:w"), "85")
                v.set(qn("w:type"), "dxa")
                mar.append(v)
            tcpr.append(mar)
            fill = "E4EBF2" if ri == 0 else ("F6F8FA" if ri % 2 == 0 else "FFFFFF")
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), fill)
            tcpr.append(shd)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.05
                for r in para.runs:
                    r.font.size = Pt(10.5)
                    r.bold = ri == 0
    p("").paragraph_format.space_after = Pt(0)
    return t


def picture(name, width, caption):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next = True
    shape = para.add_run().add_picture(str(ASSETS / name), width=Cm(width))
    shape._inline.docPr.set("descr", caption)
    cap = p(caption, "Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def reference(number, text, url=None):
    para = p(f"[{number}] {text}")
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_after = Pt(8)
    if url:
        rel = doc.part.relate_to(url,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), rel)
        r = OxmlElement("w:r")
        rp = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "1D4D75")
        rp.append(color)
        r.append(rp)
        text = OxmlElement("w:t")
        text.set(qn("xml:space"), "preserve")
        text.text = " Mở tài liệu"
        r.append(text)
        link.append(r)
        para._p.append(link)


figures()
p("BÁO CÁO LÝ THUYẾT GABOR", "Title")
p("Trích xuất đặc trưng kết cấu và tìm kiếm ảnh động vật", "Subtitle")
p("Bộ lọc Gabor mô tả biến thiên độ sáng theo hướng và tần số không gian trong một vùng ảnh. Trong hệ thống tìm kiếm ảnh động vật, đặc trưng này được dùng để biểu diễn hoa văn và cấu trúc lông. Dự án xây dựng vector Gabor 512 chiều từ ảnh xám 128 × 128, kết hợp với HSV, LBP và HOG để xếp hạng ảnh tương đồng. [R1–R4]")
h("1 Khái niệm và vai trò của Gabor")
p("Kết cấu ảnh, hay texture, thể hiện cách các mẫu sáng tối phân bố và lặp lại trên bề mặt. Những đường sọc, đốm hoặc vân lông có thể khác nhau về độ mịn, khoảng cách và hướng. Gabor kết hợp một sóng điều hòa với cửa sổ Gaussian: sóng chọn kiểu biến thiên cần quan sát, còn Gaussian giới hạn vùng ảnh tham gia phân tích. [1, 2]")
p("Một bộ lọc chỉ khảo sát một miền hướng và tần số nhất định. Khi thay đổi bước sóng và góc quay, ta có ngân hàng bộ lọc Gabor. Mỗi kênh tạo một ảnh đáp ứng, từ đó rút các thống kê gọn hơn để so sánh ảnh. Ví dụ chính thức của scikit-image dùng mean và variance của đáp ứng làm đặc trưng phân loại kết cấu. [3]")
h("Phạm vi ứng dụng trong dự án", 2)
p("Mã nguồn dùng Gabor thực, được tạo bằng OpenCV, để bổ sung thông tin kết cấu cho ảnh đã có mặt nạ alpha. Ảnh đầu vào của bộ tiền xử lý là RGBA 512 × 512; nhánh Gabor làm việc trên bản xám 128 × 128. Kết quả không phải nhãn loài mà là vector số mô tả các đáp ứng lọc trong 16 vùng không gian. [R1, R2]")
p("Báo cáo trình bày mô hình toán học, ý nghĩa tham số, quy trình tạo vector và cách tính độ tương đồng, đồng thời phân biệt các tính chất lý thuyết với lựa chọn cài đặt cụ thể trong dự án.")
h("Nội dung chính", 2)
contents = p("2 Mô hình toán học và các tham số\n3 Ngân hàng Gabor trong dự án\n4 Tiền xử lý và bảo vệ vùng kết cấu\n5 Xây dựng vector đặc trưng 512 chiều\n6 Kết hợp đặc trưng và tìm kiếm ảnh\n7 Giới hạn và cách đánh giá\n8 Tài liệu tham khảo")
contents.alignment = WD_ALIGN_PARAGRAPH.LEFT

page("2 Mô hình toán học và các tham số")
p("Với tọa độ đã quay x′ và y′, một kernel Gabor thực thường được biểu diễn bằng tích của Gaussian và hàm cosine. Công thức dưới đây dùng quy ước bước sóng λ và tỉ lệ hình học γ tương ứng các tham số của OpenCV. [1, 2]")
eq(["g(x,y) = exp(−", frac([sup("x′", "2"), " + ", sup("γ", "2"), sup("y′", "2")], ["2", sup("σ", "2")]), ") cos(", frac("2πx′", "λ"), " + ψ)"], 1)
eq(["x′ = x cos θ + y sin θ     ;     y′ = −x sin θ + y cos θ"], 2)
table(["Tham số", "Ý nghĩa", "Giá trị trong dự án"], [
    ["λ", "Bước sóng của thành phần dao động", "4, 8, 16, 32 pixel"],
    ["θ", "Góc của pháp tuyến các dải song song", "0°, 45°, 90°, 135°"],
    ["σ", "Độ lệch chuẩn Gaussian theo trục x′", "0,56λ"],
    ["γ", "Tỉ lệ hình học của bao Gaussian", "0,5"],
    ["ψ", "Độ lệch pha", "0 radian"],
    ["ksize", "Kích thước cửa sổ rời rạc", "21 × 21 pixel"],
], [2.0, 8.5, 6.2])
p("Bước sóng λ nhỏ khảo sát các biến thiên dày hơn; λ lớn khảo sát biến thiên thưa hơn. Tần số không gian f = 1/λ có đơn vị chu kỳ trên pixel. Các giá trị này được xác định trên ảnh xám 128 × 128, vì vậy không thể diễn giải trực tiếp thành kích thước thật của sọc hoặc đốm trên con vật. [R1, R2]")
p("Góc θ chỉ hướng pháp tuyến của các dải kernel, không phải hướng của chính các dải. Khi θ = 0, thành phần cosine biến đổi theo trục x và tạo các dải dọc. Tham số γ = 0,5 làm bao Gaussian kéo dài theo y′; theo biểu thức (1), độ lệch chuẩn theo y′ là σ/γ. [1]")
h("Gabor phức và trị tuyệt đối đáp ứng thực", 2)
eq([sub("g", "c"), "(x,y) = exp(−", frac([sup("x′", "2"), " + ", sup("γ", "2"), sup("y′", "2")], ["2", sup("σ", "2")]), ") exp(j", frac("2πx′", "λ"), ")"], 3)
p("Gabor phức tạo hai đáp ứng tương ứng cosine và sine. Biên độ phức được tính bằng căn tổng bình phương hai đáp ứng. Repo chỉ lọc với pha ψ = 0 rồi lấy abs, tức trị tuyệt đối của một đáp ứng thực. Hai cách này không đồng nhất; đặc trưng hiện tại vẫn có thể thay đổi theo pha và vị trí mẫu. [2, R1]")

page("3 Ngân hàng Gabor trong dự án")
p("Cấu hình gồm bốn bước sóng nhân đôi liên tiếp và bốn hướng cách đều trong nửa vòng tròn. Hai vòng lặp trong extract_gabor_spatial tạo 4 × 4 = 16 kernel. Mỗi kernel được trừ trung bình rồi chia chuẩn L2 trước khi áp dụng lên toàn bộ ảnh xám. [R1, R2]")
picture("kernel_bank.png", 14.0, "Hình 1  Ngân hàng 16 kernel theo cấu hình repo sau trừ trung bình và chuẩn hóa L2. Mỗi ô dùng thang màu đối xứng riêng để thấy hình dạng kernel.")
eq(["K₀ = K − mean(K)     ;     K̃ = ", frac("K₀", "‖K₀‖₂ + ε")], 4)
p("Trừ trung bình làm tổng hệ số kernel xấp xỉ 0, giảm đáp ứng với thành phần độ sáng không đổi. Chia chuẩn L2 cân bằng độ lớn hệ số giữa các kernel; đây là chuẩn hóa kernel, khác với chuẩn hóa vector đặc trưng ở bước cuối. Hằng số ε = 10⁻¹² giúp tránh phép chia cho 0. [R1, R2]")
p("Cửa sổ 21 × 21 được giữ cố định ở cả bốn bước sóng. Với λ = 32, σ = 17,92 pixel lớn hơn bán kính cửa sổ 10 pixel, nên bao Gaussian bị cắt đáng kể. Vì vậy các kernel thực tế là bộ lọc hữu hạn theo cấu hình repo; không nên mặc định chúng tạo phân hoạch tần số lý tưởng hoặc là cấu hình tối ưu cho mọi bộ ảnh. Đây là hệ quả trực tiếp của các tham số trong mã. [R1, R2]")

page("4 Tiền xử lý và bảo vệ vùng kết cấu")
h("4 1 Chuẩn bị ảnh xám và mặt nạ", 2)
p("Hàm prepare_image nhận mảng uint8 RGBA 512 × 512, tạo mặt nạ bằng alpha > 0 và đặt RGB ngoài mặt nạ về 0. Ảnh được đổi sang xám, thu nhỏ về 128 × 128 bằng INTER_AREA; mặt nạ dùng INTER_NEAREST. Sau khi gán lại nền xám bằng 0, nhánh Gabor đổi ảnh sang float32 và chia 255. Alpha rỗng gây lỗi; alpha dương được xét như foreground, không gán trọng số theo độ trong suốt. [R1, R2]")
h("4 2 Loại ảnh hưởng của biên đối tượng", 2)
p("Biên giữa con vật và nền xám 0 có thể tạo đáp ứng mạnh. Để ưu tiên hoa văn bên trong lông, repo co mặt nạ bằng phần tử cấu trúc vuông 21 × 21 với giá trị ngoài ảnh bằng 0. Một tâm chỉ được giữ khi toàn cửa sổ nằm trong foreground. [R1]")
p("Nếu mặt nạ sau co 21 × 21 rỗng trên toàn ảnh, chương trình chuyển sang co 3 × 3. Nhánh dự phòng giữ thêm vị trí để thống kê, nhưng không bảo đảm cửa sổ Gabor 21 × 21 tránh được nền. Nếu kết quả dự phòng cũng rỗng, vector có norm bằng 0 và bị đánh dấu không hợp lệ. Việc dự phòng áp dụng cho toàn ảnh, không thực hiện riêng cho từng ô lưới. [R1]")
h("4 3 Lọc và lấy độ lớn đáp ứng", 2)
eq([sub("A", "s,t"), "(x,y) = |filter2D(I, ", sub("K̃", "s,t"), ")(x,y)|"], 5)
p("Chỉ số s biểu thị bước sóng và t biểu thị hướng. Toàn ảnh được lọc trước; mặt nạ chỉ quyết định vị trí lấy thống kê sau lọc. Trị tuyệt đối tránh triệt tiêu các giá trị dương và âm khi lấy trung bình. Trong báo cáo, A được gọi chính xác là độ lớn đáp ứng thực; mã không tính năng lượng bình phương |R|². [R1]")
p("OpenCV filter2D thực hiện tương quan, không lật kernel như tích chập. Với pha 0 và kernel đối xứng tâm đang dùng, hai phép cho kết quả tương đương về toán học, bỏ qua sai số số học. Khi thay đổi kernel hoặc pha, cần kiểm tra lại quy ước này. [1, R1]")
picture("orientation_demo.png", 14.5, "Hình 2  Minh họa vân tổng hợp chu kỳ 8 pixel bằng kernel của repo. Các ảnh dùng cùng thang xám. Hình minh họa nguyên lý, không phải đánh giá kho ảnh.")

page("5 Xây dựng vector đặc trưng 512 chiều")
h("5 1 Thống kê theo vùng", 2)
p("Ảnh xám 128 × 128 được chia thành 4 hàng và 4 cột, tạo 16 ô kích thước 32 × 32. Trong mỗi ô và mỗi kênh Gabor, chương trình chọn pixel có tâm hợp lệ sau co mặt nạ, rồi tính mean và độ lệch chuẩn tổng thể. [R1]")
eq([sub("μ", "r,s,t"), " = ", frac("1", sub("N", "r")), " ∑ ", sub("A", "s,t"), "(x,y)"], 6)
eq([sub("σ", "r,s,t"), " = ", root([frac("1", sub("N", "r")), " ∑ ", sup(["(", sub("A", "s,t"), "(x,y) − ", sub("μ", "r,s,t"), ")"], "2")])], 7)
p("Trong (6) và (7), tổng lấy trên các tâm hợp lệ thuộc ô r, Nᵣ là số tâm đó. Hàm numpy.std mặc định ddof = 0 nên mẫu số là Nᵣ. Mean biểu thị mức đáp ứng trung bình, còn độ lệch chuẩn biểu thị mức phân tán trong vùng. Ô không có tâm hợp lệ nhận cặp (0, 0) cho mọi kênh và được ghi trong empty_regions. [R1]")
h("5 2 Thứ tự ghép và số chiều", 2)
p("Thứ tự ghép cố định là hàng ô, cột ô, bước sóng tăng dần, hướng tăng dần, rồi mean và std. Một ô có 16 kênh × 2 thống kê = 32 giá trị; cả ảnh có 16 ô × 32 = 512 giá trị. Hai ảnh chỉ so sánh được khi dùng cùng cấu hình, quy trình tiền xử lý và thứ tự ghép. [R1, R2]")
eq([sub("d", "Gabor"), " = 4 × 4 × 4 × 4 × 2 = 512"], 8)
p("Các vị trí 0 và 1 lần lượt là mean và std của ô trên cùng bên trái, λ = 4 và θ = 0°. Vị trí 2 và 3 thuộc cùng ô, λ = 4 và θ = 45°. Sau 32 giá trị của một ô, vector chuyển sang ô bên phải. Với r và c bắt đầu từ 0, kênh q = 4s + t, chỉ số mean là 32(4r + c) + 2q; chỉ số std lớn hơn 1. [R1]")
h("5 3 Chuẩn hóa L2", 2)
eq([sub("f̂", "Gabor"), " = ", frac(sub("f", "Gabor"), "‖fGabor‖₂"), "     khi ‖fGabor‖₂ > ε"], 9)
p("Hàm _finalize kiểm tra vector đủ 512 chiều, hữu hạn và không âm. Nếu norm lớn hơn ε, vector được chia cho norm; ngược lại nhận vector 0 và valid = False. Chuẩn hóa giảm ảnh hưởng của một hệ số nhân chung lên toàn vector, nhưng không bảo đảm bất biến trước bóng đổ, bão hòa sáng hoặc thay đổi mặt nạ. [R1]")
p("Lưới giúp giữ thông tin vị trí tương đối của kết cấu. Đổi lại, sự dịch chuyển hoặc thay đổi tư thế có thể đưa cùng hoa văn sang ô khác, làm vector thay đổi. Đây là đặc điểm của cách thống kê không gian hiện tại, cần cân nhắc khi diễn giải kết quả. [R1]")

page("6 Kết hợp đặc trưng và tìm kiếm ảnh")
p("Gabor là một trong năm nhóm đặc trưng của pipeline. Mỗi nhóm được chuẩn hóa L2 riêng, sau đó nhân căn bậc hai trọng số và ghép thành vector chung. Cấu hình hiện tại ưu tiên Gabor với trọng số 0,50; đây là lựa chọn của dự án, không phải hệ số suy ra từ lý thuyết Gabor. [R2, R3]")
table(["Nhóm đặc trưng", "Số chiều", "Trọng số"], [
    ["HSV toàn cục", 288, "0,05"], ["HSV lưới 2 × 2", 1152, "0,15"],
    ["LBP lưới 4 × 4", 160, "0,15"], ["HOG", 1764, "0,15"],
    ["Gabor lưới 4 × 4", 512, "0,50"], ["Tổng", 3876, "1,00"],
], [10.0, 3.35, 3.35])
eq(["F = [", root(sub("w", "1")), sub("f̂", "1"), "; … ; ", root(sub("w", "5")), sub("f̂", "5"), "]"], 10)
p("Khi các nhóm hợp lệ và tổng trọng số bằng 1, vector F có norm xấp xỉ 1. Tích vô hướng hai vector F bằng tổng có trọng số của năm tích vô hướng thành phần. Pipeline chỉ tạo F khi cả năm nhóm hợp lệ. [R3]")
h("6 1 Độ tương đồng thành phần", 2)
eq([sub("s", "cos"), "(q,d) = ", sub("f̂", "q"), " · ", sub("f̂", "d")], 11)
p("Ở chế độ cosine, chỉ mục tính tích vô hướng giữa các vector thành phần đã chuẩn hóa rồi chặn kết quả trong [0, 1]. Với Gabor, các thành phần không âm nên độ tương đồng này không âm về lý thuyết. Ở chế độ Euclidean, chương trình tính khoảng cách L2 rồi đổi thành điểm 1/(1 + d). [R4]")
eq(["S(q,d) = 0,05", sub("s", "HSV"), " + 0,15", sub("s", "HSV2×2"), " + 0,15", sub("s", "LBP"), " + 0,15", sub("s", "HOG"), " + 0,50", sub("s", "Gabor")], 12)
h("6 2 Vai trò của PCA", 2)
p("Chỉ mục fit PCA trên vector ghép 3.876 chiều và giữ tối đa 128 thành phần, đồng thời không vượt số ảnh trừ một hoặc hạng số có ý nghĩa của dữ liệu. Mã tính tổng điểm năm nhóm cho tất cả ảnh và dùng tổng điểm làm khóa xếp hạng chính. Điểm PCA chỉ là khóa phụ khi tổng điểm bằng nhau; mã hiện tại không truy hồi ứng viên bằng PCA trước rồi mới rerank. [R4]")
p("Điểm Gabor cao cho biết các thống kê đáp ứng gần nhau. Nó không phải xác suất cùng loài, không trực tiếp đếm đốm và không đủ để kết luận hai ảnh có cùng hình dạng hoa văn. [R1, R4]")

page("7 Giới hạn và cách đánh giá")
h("7 1 Giới hạn của biểu diễn hiện tại", 2)
p("Bộ lọc đa hướng không tự tạo ra bất biến quay. Khi ảnh quay, năng lượng có thể chuyển giữa các kênh và các ô lưới. Repo chưa căn chỉnh hướng trội hoặc hoán vị kênh để xử lý thay đổi này. Tổng quan về truy hồi kết cấu cũng phân biệt Gabor thông thường với các mở rộng bất biến quay. [4, R1]")
p("Bước sóng cố định chỉ có ý nghĩa trên ảnh đã chuẩn bị. Nếu đối tượng chiếm tỉ lệ khác nhau trên canvas hoặc ảnh bị mờ khi thu nhỏ, cấu trúc lông quan sát được có thể khác. Kernel 21 × 21 còn cắt mạnh các thang lớn; mở rộng cửa sổ để giữ bao Gaussian sẽ đồng thời làm vùng mặt nạ hợp lệ sau co nhỏ đi. Hai lựa chọn cần được đánh giá cùng nhau. [R1, R2]")
p("Mean và std chỉ tóm tắt hai thống kê của độ lớn đáp ứng. Những bề mặt khác nhau vẫn có thể tạo giá trị gần nhau; các quan hệ vị trí chi tiết trong từng ô bị mất. Chất lượng tách nền, diện tích foreground và số ô rỗng cũng ảnh hưởng vector. Khi dùng nhánh co 3 × 3 dự phòng, biên đối tượng có thể quay lại đóng góp đáng kể. [R1]")
h("7 2 Đánh giá phù hợp với bài toán", 2)
p("Các kiểm tra về số chiều, giá trị hữu hạn và norm L2 xác nhận tính hợp lệ kỹ thuật của bộ trích xuất. Chúng không đo chất lượng truy hồi hay chứng minh trọng số 0,50 tối ưu. Để đánh giá, cần xác định rõ ảnh liên quan là cùng loài hay cùng hoa văn, rồi dùng một tập truy vấn có nhãn. [R2, R5]")
eq(["Precision@K = ", frac("Số ảnh liên quan trong K kết quả đầu", "K")], 13)
p("Nên so sánh ít nhất ba cấu hình: chỉ Gabor, kết hợp không có Gabor và kết hợp đầy đủ. Với mỗi cấu hình, báo cáo Precision@K, thời gian trích xuất, thời gian tìm kiếm và tỉ lệ vector không hợp lệ. Khi thay đổi trọng số hoặc kernel, chọn tham số trên tập kiểm định rồi đánh giá trên tập kiểm thử độc lập. Loại chính ảnh truy vấn và các bản sao gần trùng khỏi tập ứng viên để tránh làm điểm số quá thuận lợi.")
p("Thử nghiệm bổ sung có thể thay đổi góc quay, tỉ lệ đối tượng, độ sáng và chất lượng mặt nạ để xác định nguyên nhân sai lệch. Các phép thử này là đề xuất đánh giá, chưa phải kết quả thực nghiệm của báo cáo.")
h("Kết luận", 2)
p("Thiết kế hiện tại biểu diễn kết cấu bằng 16 kênh Gabor thực, thống kê mean và std trên lưới 4 × 4 rồi chuẩn hóa thành vector 512 chiều. Cách biểu diễn có thể giải thích theo hướng, bước sóng và vùng ảnh; hiệu quả nhận biết hoa văn cần được xác nhận bằng đánh giá truy hồi có nhãn, đặc biệt khi Gabor chiếm một nửa tổng điểm. [R1–R4]")

page("8 Tài liệu tham khảo")
p("Các tài liệu trực tuyến dưới đây có nội dung đọc miễn phí. Thông tin về cách triển khai được đối chiếu trực tiếp với các tệp mã nguồn liệt kê ở cuối mục. Ngày truy cập tài liệu: 21 tháng 10 năm 2026.")
reference(1, "OpenCV. Image Filtering. Tài liệu các hàm getGaborKernel, filter2D và erode. Dùng để đối chiếu ý nghĩa tham số và quy ước lọc.", "https://docs.opencv.org/4.x/d4/d86/group__imgproc__filter.html")
reference(2, "scikit-image. skimage.filters.gabor và gabor_kernel. Tài liệu bộ lọc thực, ảo và kernel Gabor phức.", "https://scikit-image.org/docs/stable/api/skimage.filters.html#skimage.filters.gabor")
reference(3, "scikit-image. Gabor filter banks for texture classification. Ví dụ mã nguồn về ngân hàng lọc và thống kê mean, variance.", "https://scikit-image.org/docs/stable/auto_examples/features_detection/plot_gabor.html")
reference(4, "N. Baaziz, O. Abahmane và R. Missaoui. Texture feature extraction in the spatial-frequency domain for content-based image retrieval. arXiv:1012.5208, 2010. Xem mục 4 về Gabor và bất biến quay.", "https://arxiv.org/abs/1012.5208")
h("Nguồn mã trong dự án", 2)
for code, text in [
    ("R1", "app/search/extractors.py — hàm extract_gabor_spatial và _finalize; app/search/preprocess.py — hàm prepare_image."),
    ("R2", "app/search/config.py — lớp FeatureConfig, FEATURE_DIMS và trọng số các nhóm."),
    ("R3", "app/search/pipeline.py — hàm extract_features và cách ghép vector theo căn trọng số."),
    ("R4", "app/search/index.py — LocalSearchIndex, PCA và thứ tự xếp hạng."),
    ("R5", "tests/test_search.py — kiểm tra số chiều, tính hữu hạn và chuẩn hóa của vector Gabor."),
]:
    p(f"[{code}] {text}").alignment = WD_ALIGN_PARAGRAPH.LEFT
h("Đối chiếu các công thức với mã nguồn", 2)
p("Công thức (1) đến (3) mô tả mô hình bộ lọc và phân biệt dạng thực với dạng phức. Công thức (4) đến (12) diễn giải các phép tính đang có trong repo. Công thức (13) là chỉ số đề xuất cho đánh giá truy hồi; báo cáo không công bố độ chính xác hoặc tốc độ chưa đo được.")

doc.core_properties.title = "Lý thuyết Gabor trong trích xuất đặc trưng và tìm kiếm ảnh động vật"
doc.core_properties.subject = "Bộ lọc Gabor đa tỉ lệ đa hướng và vector không gian 512 chiều"
doc.core_properties.author = ""
doc.core_properties.keywords = "Gabor, texture, CBIR, OpenCV, 512 chiều"
doc.save(OUT)
assert CFG.gabor_rows * CFG.gabor_cols * len(CFG.gabor_wavelengths) * CFG.gabor_orientations * 2 == 512
print(OUT)
print(f"Paragraphs: {len(doc.paragraphs)}; tables: {len(doc.tables)}; figures: {len(doc.inline_shapes)}")
