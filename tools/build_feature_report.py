from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "output" / "report_assets"
OUT = ROOT / "output" / "documents" / "Bao_cao_trich_xuat_dac_trung_Animal_Image_Search.docx"
DATA = json.loads((ASSETS / "run_summary.json").read_text(encoding="utf-8"))
DETAILS = DATA["details"]

INK = "111827"
MUTED = "475569"
NAVY = "163A5F"
BLUE = "2563EB"
PALE = "EEF4FA"
GRID = "D9D9D9"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = GRID, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def clear_paragraph_borders(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Arial", size=None, bold=None, color=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr_text, fld_end])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(6)

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(16)
    title_ppr = title._element.get_or_add_pPr()
    title_borders = title_ppr.find(qn("w:pBdr"))
    if title_borders is not None:
        title_ppr.remove(title_borders)

    for name, size, before, after in (("Heading 1", 16, 14, 8), ("Heading 2", 13, 12, 6), ("Heading 3", 11, 10, 4)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run("ANIMAL IMAGE SEARCH  |  BÁO CÁO ĐẶC TRƯNG")
    set_run_font(r, size=8, bold=True, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Trang ")
    set_run_font(r, size=8, color=MUTED)
    add_field(fp, "PAGE")


def add_heading(doc, text: str, level: int = 1):
    return doc.add_heading(text, level=level)


def add_body(doc, text: str, *, bold_lead: str | None = None):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_bullets(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_run_font(p.add_run(item))
        p.paragraph_format.space_after = Pt(3)


def add_table(doc, headers: list[str], rows: list[list[str]], widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, size=9, bold=True, color="FFFFFF")
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cell = cells[idx]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, size=9)
            if row_index % 2:
                set_cell_shading(cell, PALE)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(3)
    return table


def set_picture_alt(inline_shape, description: str) -> None:
    inline_shape._inline.docPr.set("descr", description)


def add_figure(doc, filename: str, caption: str, width=6.25) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(ASSETS / filename), width=Inches(width))
    set_picture_alt(shape, caption)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = False
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    set_run_font(r, size=9, color=MUTED)
    r.italic = True


def add_formula(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text)
    set_run_font(r, name="Cambria Math", size=11)


def build() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(54)
    r = p.add_run("BÁO CÁO KỸ THUẬT")
    set_run_font(r, size=11, bold=True, color=BLUE)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Trích xuất đặc trưng và tìm kiếm ảnh trong Animal Image Search")
    clear_paragraph_borders(title)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    r = subtitle.add_run("Phân tích trực tiếp từ mã nguồn và kết quả chạy repository")
    set_run_font(r, size=12, color=MUTED)

    cover_shape = doc.add_picture(
        str(ASSETS / "03_ket_qua_tim_kiem.png"), width=Inches(6.3)
    )
    set_picture_alt(
        cover_shape,
        "Ảnh truy vấn tiger_001.png và năm kết quả tương đồng do repository trả về",
    )
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.paragraph_format.space_before = Pt(22)
    r = info.add_run("Ngày lập báo cáo 19 tháng 10 năm 2026")
    set_run_font(r, size=10, color=MUTED)
    doc.add_page_break()

    add_heading(doc, "Tóm tắt kết quả", 1)
    add_body(
        doc,
        "Repository triển khai đúng bốn nhóm đặc trưng thủ công cho ảnh động vật đã tách nền: histogram HSV toàn cục, histogram HSV theo lưới 2×2, LBP theo lưới 4×4 và HOG. Bốn vector có số chiều lần lượt là 288, 1.152, 160 và 1.764. Sau chuẩn hóa L2, chúng được nhân căn bậc hai của trọng số rồi ghép thành vector 3.364 chiều. HOG nhận trọng số lớn nhất 0,35; hai đặc trưng không gian HSV và LBP cùng nhận 0,25; HSV toàn cục nhận 0,15.",
    )
    add_body(
        doc,
        f"Lần chạy minh họa trên {DETAILS['index']['images']} ảnh hợp lệ tạo PCA {DETAILS['pca_dimension']} chiều. Truy vấn tiger_001.png dùng trực tiếp alpha sẵn có và hoàn thành trong {DETAILS['elapsed_seconds']:.3f} giây, bao gồm thời gian lập chỉ mục lần đầu. Ảnh truy vấn đứng hạng 1 với điểm 1,000000; bốn vị trí tiếp theo đều là ảnh hổ. Kết quả này xác nhận đường đi từ ảnh RGBA qua bốn bộ trích xuất, ghép vector, PCA và xếp hạng hoạt động đồng bộ.",
    )

    add_heading(doc, "Phạm vi báo cáo", 1)
    add_body(
        doc,
        "Báo cáo chỉ mô tả các đặc trưng được gọi trong app/search/extractors.py và được ghép trong app/search/pipeline.py. Phần YOLO World và BiRefNet chỉ được nhắc ở mức tạo ảnh RGBA đầu vào cho trích xuất; chúng không phải vector đặc trưng tìm kiếm. Các phương pháp không xuất hiện trong pipeline tìm kiếm không được đưa vào nội dung.",
    )
    add_table(
        doc,
        ["Nhóm", "Kích thước đầu vào", "Cấu hình chính", "Số chiều", "Trọng số"],
        [
            ["HSV toàn cục", "RGB 512×512 và alpha", "18×4×4 bins", "288", "0,15"],
            ["HSV lưới", "RGB 512×512 và alpha", "Lưới 2×2; 288 bins mỗi vùng", "1.152", "0,25"],
            ["LBP lưới", "Gray 128×128 và mask", "P=8, R=1, uniform; lưới 4×4", "160", "0,25"],
            ["HOG", "Gray 128×128", "9 hướng; cell 16×16; block 2×2", "1.764", "0,35"],
        ],
        [2.6, 3.2, 6.2, 2.0, 1.8],
    )

    add_heading(doc, "Chuẩn hóa ảnh trước trích xuất", 1)
    add_body(
        doc,
        "Ảnh tìm kiếm phải là RGBA 512×512 và có ít nhất một pixel alpha lớn hơn 0. Nếu ảnh có kích thước khác, hàm standardize_rgba căn giữa đối tượng trên canvas 512×512 trong khi giữ tỷ lệ. Hàm prepare_image tạo mask nhị phân từ điều kiện alpha > 0, đặt RGB ngoài mask về 0, chuyển RGB sang ảnh xám và thu nhỏ ảnh xám cùng mask xuống 128×128. Ảnh xám dùng nội suy vùng, còn mask dùng lân cận gần nhất để không tạo giá trị mask trung gian.",
    )
    add_figure(doc, "01_tien_xu_ly.png", "Hình 1  Dữ liệu đầu vào thực tế của các bộ trích xuất")

    add_heading(doc, "Histogram HSV toàn cục", 1)
    add_body(
        doc,
        "Ảnh RGB 512×512 được đổi sang HSV theo quy ước OpenCV, trong đó H nằm trong khoảng 0 đến 179 và S, V nằm trong khoảng 0 đến 255. Mã nguồn chia H thành 18 khoảng, S thành 4 khoảng và V thành 4 khoảng. Những pixel có S < 16 hoặc V < 16 được gán H = 0 để giảm dao động sắc độ ở vùng gần xám hoặc quá tối. Chỉ pixel nằm trong alpha mask mới được đếm.",
    )
    add_formula(doc, "dHSV toàn cục = 18 × 4 × 4 = 288")
    add_body(
        doc,
        "Số đếm từng bin được chia cho số pixel foreground, sau đó toàn vector được chuẩn hóa L2. Đặc trưng này mô tả phân bố màu tổng thể của con vật nhưng không giữ vị trí xuất hiện của màu trong ảnh.",
    )

    add_heading(doc, "Histogram HSV theo lưới 2×2", 1)
    add_body(
        doc,
        "Bộ trích xuất thứ hai dùng cùng phép lượng tử HSV, nhưng chia ảnh thành bốn vùng bằng nhau theo hai hàng và hai cột. Mỗi vùng sinh một histogram 288 chiều chỉ từ pixel foreground trong vùng đó. Bốn histogram được ghép theo thứ tự trái sang phải và từ trên xuống dưới. Vùng không có foreground tạo vector 0 thay vì gây lỗi.",
    )
    add_formula(doc, "dHSV không gian = 2 × 2 × 288 = 1.152")
    add_body(
        doc,
        "Việc giữ bố cục thô giúp phân biệt những ảnh có bảng màu gần nhau nhưng màu phân bố ở các phần cơ thể khác nhau. Sau khi ghép, toàn vector 1.152 chiều được chuẩn hóa L2 một lần.",
    )

    add_heading(doc, "LBP theo lưới 4×4", 1)
    add_body(
        doc,
        "Local Binary Pattern được tính trên ảnh xám 128×128 với P=8 điểm lân cận, bán kính R=1 và phương pháp uniform của scikit image. Mask 128×128 được co một lần bằng kernel 3×3 trước khi thống kê, vì tâm LBP sát biên alpha có thể dùng lân cận thuộc nền. Bước co mask loại các tâm không đủ vùng lân cận foreground.",
    )
    add_body(
        doc,
        "Ảnh được chia thành lưới 4×4. Với cấu hình uniform P=8, mỗi vùng dùng 10 bin. Histogram từng vùng được chuẩn hóa theo số tâm hợp lệ trong vùng, sau đó 16 histogram được ghép và chuẩn hóa L2.",
    )
    add_formula(doc, "dLBP = 4 × 4 × 10 = 160")
    add_body(
        doc,
        "LBP bổ sung thông tin vi cấu trúc sáng tối, phù hợp với họa tiết lông, vằn và đốm mà histogram màu không thể biểu diễn trực tiếp.",
    )

    add_heading(doc, "Histogram hướng gradient HOG", 1)
    add_body(
        doc,
        "HOG cũng dùng ảnh xám 128×128. Cấu hình có 9 hướng gradient, cell 16×16 pixel, block 2×2 cell, chuẩn hóa block L2 Hys và không dùng transform_sqrt. Ảnh 128×128 tạo lưới 8×8 cell; cửa sổ block trượt một cell nên có 7×7 block. Mỗi block chứa 2×2×9 giá trị.",
    )
    add_formula(doc, "dHOG = 7 × 7 × 2 × 2 × 9 = 1.764")
    add_body(
        doc,
        "Tensor HOG được làm phẳng theo đúng thứ tự mà thư viện trả về, rồi chuẩn hóa L2 ở cấp toàn vector. Trong pipeline hiện tại, HOG nhận trọng số cao nhất vì cung cấp thông tin hình dáng và biên định hướng của đối tượng.",
    )

    add_figure(doc, "02_vector_dac_trung.png", "Hình 2  Bốn vector L2 trích xuất từ tiger_001.png", width=6.3)

    add_heading(doc, "Ghép vector và chuẩn hóa trọng số", 1)
    add_body(
        doc,
        "Chỉ khi cả bốn vector hợp lệ, pipeline mới tạo vector kết hợp. Với mỗi nhóm i, vector L2 fi được nhân với căn bậc hai của trọng số wi rồi ghép nối. Cách nhân căn bậc hai làm cho tích vô hướng giữa hai vector kết hợp bằng tổng có trọng số của bốn độ tương đồng cosine thành phần.",
    )
    add_formula(doc, "F = [√0,15 fHSV; √0,25 fHSV2×2; √0,25 fLBP; √0,35 fHOG]")
    add_formula(doc, "dF = 288 + 1.152 + 160 + 1.764 = 3.364")
    add_body(
        doc,
        "Vì từng fi có norm bằng 1 và tổng trọng số bằng 1, norm của F cũng bằng 1 trong trường hợp mọi nhóm hợp lệ. Mã nguồn kiểm tra kích thước, tính hữu hạn và giá trị không âm trước khi cho vector tham gia tìm kiếm.",
    )

    add_heading(doc, "PCA và cách xếp hạng", 1)
    add_body(
        doc,
        "Chỉ mục cục bộ fit PCA trực tiếp trên các vector 3.364 chiều. Để tránh phân rã ma trận rất rộng, mã nguồn tính ma trận Gram theo số ảnh, phân rã trị riêng rồi dựng các trục thành phần chính. Số chiều tối đa là 128, đồng thời không vượt quá số ảnh trừ một và số trị riêng có phương sai đáng kể. Lần chạy minh họa đạt đủ 128 chiều.",
    )
    add_body(
        doc,
        "Điểm xếp hạng chính không lấy trực tiếp từ PCA. Với cosine, hệ thống tính độ tương đồng riêng cho từng nhóm đã chuẩn hóa L2 rồi cộng theo trọng số 0,15, 0,25, 0,25 và 0,35. Với khoảng cách Euclid, mỗi khoảng cách thành phần d được đổi thành 1/(1+d) trước khi cộng trọng số. PCA cung cấp pca_similarity và được dùng làm khóa phụ khi các điểm tổng hợp bằng nhau.",
    )
    add_formula(doc, "score = 0,15 sHSV + 0,25 sHSV2×2 + 0,25 sLBP + 0,35 sHOG")

    add_heading(doc, "Kết quả chạy repository", 1)
    add_body(
        doc,
        f"Kho ảnh thực tế có {DETAILS['index']['images']} ảnh hợp lệ: 67 ảnh cheetah, 64 ảnh trong nhóm lion remove, 23 ảnh trong nhóm snow_leopard remove và 63 ảnh tiger. Truy vấn dùng cosine và Top K bằng 5. Vì tiger_001.png đã có alpha hữu ích, hệ thống dùng nhánh existing_alpha, chuẩn hóa về 512×512 rồi trích xuất trực tiếp.",
    )
    rows = []
    for match in DETAILS["matches"]:
        rows.append(
            [
                str(match["rank"]),
                Path(match["source"]).name,
                match["label"],
                f"{match['score']:.6f}",
                f"{match['pca_similarity']:.6f}",
            ]
        )
    add_table(doc, ["Hạng", "Ảnh", "Nhóm", "Điểm tổng hợp", "PCA similarity"], rows, [1.4, 4.2, 3.0, 3.0, 3.0])
    add_figure(doc, "03_ket_qua_tim_kiem.png", "Hình 3  Ảnh truy vấn và Top 5 kết quả do repository trả về", width=6.25)
    add_figure(doc, "04_diem_thanh_phan.png", "Hình 4  Mức đóng góp của từng nhóm đặc trưng vào điểm tổng hợp", width=6.25)

    add_heading(doc, "Nhận xét từ kết quả", 1)
    add_bullets(
        doc,
        [
            "Ảnh truy vấn xuất hiện ở hạng 1 với điểm 1,000000 cho cả bốn thành phần, phù hợp với kiểm thử self query của chỉ mục.",
            "Bốn kết quả tiếp theo đều thuộc nhóm tiger. Điểm LBP của chúng từ 0,957137 đến 0,987077, cao hơn phần lớn điểm màu; họa tiết cục bộ góp phần ổn định nhận dạng các ảnh hổ.",
            "HOG của các hạng 3 đến 5 nằm trong khoảng 0,868596 đến 0,887257 và có trọng số 0,35, nên hình dáng và hướng biên ảnh hưởng lớn tới thứ hạng cuối.",
            "PCA similarity giảm từ 0,718909 ở hạng 2 xuống 0,669061 ở hạng 5, nhưng thứ tự chính vẫn do tổng điểm bốn nhóm quyết định.",
        ],
    )

    add_heading(doc, "Đảm bảo tính nhất quán", 1)
    add_body(
        doc,
        "Mã nguồn cố định thứ tự ghép là hsv_global, hsv_spatial, lbp_spatial, hog và lưu fingerprint của cấu hình. Bộ kiểm thử xác nhận đúng bốn kích thước, vector kết hợp 3.364 chiều có norm xấp xỉ 1, biểu đồ đọc đúng vector L2, dữ liệu NaN hoặc Infinity bị từ chối, và truy vấn chính nó đứng hạng đầu. Khi chạy toàn bộ repository trong lần lập báo cáo, 25 trên 25 kiểm thử đều đạt.",
    )

    add_heading(doc, "Kết luận", 1)
    add_body(
        doc,
        "Pipeline tìm kiếm của Animal Image Search kết hợp màu sắc, bố cục màu, vi cấu trúc và hình dạng trong một biểu diễn 3.364 chiều. Thiết kế dùng alpha mask để hạn chế nền, L2 để đưa các nhóm về cùng thang đo và trọng số để ưu tiên HOG cùng hai đặc trưng không gian. PCA 128 chiều hỗ trợ biểu diễn chỉ mục và khóa phụ, còn điểm xếp hạng cuối vẫn giữ khả năng giải thích qua bốn điểm thành phần.",
    )

    doc.core_properties.title = "Trích xuất đặc trưng và tìm kiếm ảnh trong Animal Image Search"
    doc.core_properties.subject = "Báo cáo kỹ thuật dựa trên mã nguồn và kết quả chạy repository"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "HSV, LBP, HOG, PCA, tìm kiếm ảnh"
    doc.save(OUT)
    print("output/documents/Bao_cao_trich_xuat_dac_trung_Animal_Image_Search.docx")


if __name__ == "__main__":
    build()
