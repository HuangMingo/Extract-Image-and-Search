"""Create section 3.2.2 from the current feature configuration."""
from __future__ import annotations

import sys
from pathlib import Path
from PIL import Image  # Use the installed Pillow compatible with this interpreter.

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / ".test-deps"))
sys.path.insert(0, str(ROOT))
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from app.search.config import DEFAULT_CONFIG, FEATURE_DIMS, FEATURE_ORDER, COMBINED_FEATURE_DIM

OUT = ROOT / "output/documents/Muc_3_2_2_Trich_rut_dac_trung.docx"
QA = ROOT / "tmp/feature_section_322"
QA.mkdir(parents=True, exist_ok=True)
OUT.parent.mkdir(parents=True, exist_ok=True)
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(2.1), Cm(2.0)
section.left_margin, section.right_margin = Cm(2.5), Cm(2.0)

for style in doc.styles:
    if style._element.pPr is not None:
        for item in list(style._element.pPr.findall(qn("w:pBdr"))):
            style._element.pPr.remove(item)
    if style._element.rPr is not None:
        fonts = style._element.rPr.find(qn("w:rFonts"))
        if fonts is not None:
            for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                fonts.attrib.pop(qn("w:" + attr), None)
for name in ("Normal", "Title", "Caption"):
    style = doc.styles[name]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = False
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal = doc.styles["Normal"].paragraph_format
normal.line_spacing = 1.3
normal.space_after = Pt(6)
normal.first_line_indent = Cm(0.75)
normal.widow_control = True
title_style = doc.styles["Title"]
title_style.font.size = Pt(14)
title_style.font.color.rgb = RGBColor.from_string("155E7C")
title_style.paragraph_format.space_before = Pt(0)
title_style.paragraph_format.space_after = Pt(9)
title_style.paragraph_format.keep_with_next = True
title_style.paragraph_format.first_line_indent = Cm(0.75)
caption_style = doc.styles["Caption"]
caption_style.font.italic = True
caption_style.font.color.rgb = RGBColor.from_string("17365D")
caption_style.paragraph_format.first_line_indent = Cm(0)
caption_style.paragraph_format.space_after = Pt(8)


def p(text, style=None):
    para = doc.add_paragraph(text, style)
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return para


def bullet(text):
    para = p("–  " + text)
    para.paragraph_format.left_indent = Cm(1.0)
    para.paragraph_format.first_line_indent = Cm(-0.45)
    para.paragraph_format.space_after = Pt(2)
    return para


def run(text):
    r = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def add(target, values):
    for value in values if isinstance(values, (list, tuple)) else [values]:
        target.append(run(value) if isinstance(value, str) else value)


def op(tag, fields):
    element = OxmlElement("m:" + tag)
    for name, value in fields:
        part = OxmlElement("m:" + name)
        add(part, value)
        element.append(part)
    return element


def sub(base, index):
    return op("sSub", [("e", base), ("sub", index)])


def hat(value):
    element = OxmlElement("m:acc")
    pr = OxmlElement("m:accPr")
    character = OxmlElement("m:chr")
    character.set(qn("m:val"), "̂")
    pr.append(character)
    element.append(pr)
    expression = OxmlElement("m:e")
    add(expression, value)
    element.append(expression)
    return element


def sup(base, power):
    return op("sSup", [("e", base), ("sup", power)])


def frac(num, den):
    return op("f", [("num", num), ("den", den)])


def root(value):
    element = OxmlElement("m:rad")
    properties = OxmlElement("m:radPr")
    hidden = OxmlElement("m:degHide")
    hidden.set(qn("m:val"), "1")
    properties.append(hidden)
    element.append(properties)
    element.append(OxmlElement("m:deg"))
    radicand = OxmlElement("m:e")
    add(radicand, value)
    element.append(radicand)
    return element


def summation(lower, upper, expression):
    element = OxmlElement("m:nary")
    pr = OxmlElement("m:naryPr")
    character = OxmlElement("m:chr")
    character.set(qn("m:val"), "∑")
    pr.append(character)
    limit = OxmlElement("m:limLoc")
    limit.set(qn("m:val"), "subSup")
    pr.append(limit)
    element.append(pr)
    for name, value in (("sub", lower), ("sup", upper), ("e", expression)):
        part = OxmlElement("m:" + name)
        add(part, value)
        element.append(part)
    return element


def eq(values):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.first_line_indent = Cm(0)
    para.paragraph_format.space_before = Pt(5)
    para.paragraph_format.space_after = Pt(9)
    math = OxmlElement("m:oMath")
    add(math, values)
    para._p.append(math)
    return para


p("3.2.2. Thực hiện trích rút các đặc trưng đã chọn từ bộ dữ liệu", "Title")
p("Sau khi bộ dữ liệu được tiền xử lý, đối tượng chính được làm nổi bật và ảnh được chuẩn hóa về dạng RGBA 512 × 512, hệ thống thực hiện trích rút đặc trưng theo cùng một quy trình cho từng ảnh. Mặt nạ alpha xác định vùng đối tượng; các giá trị RGB ngoài vùng này được đặt bằng 0. Đặc trưng màu được tính trên ảnh 512 × 512, còn các đặc trưng LBP, HOG và Gabor sử dụng ảnh xám 128 × 128.")
p("Với mỗi ảnh, hệ thống tạo ra năm nhóm mô tả bổ trợ nhau:")
bullet("Đặc trưng phân bố màu toàn cục bằng histogram HSV (HSV Global).")
bullet("Đặc trưng phân bố màu theo không gian bằng histogram HSV trên lưới 2 × 2 (HSV Spatial).")
bullet("Đặc trưng kết cấu cục bộ bằng LBP trên lưới 4 × 4 (LBP Spatial).")
bullet("Đặc trưng hướng gradient và cấu trúc biên bằng HOG.")
bullet("Đặc trưng kết cấu đa tỉ lệ, đa hướng bằng Gabor trên lưới 4 × 4 (Gabor Spatial).")
p("Histogram HSV toàn cục gồm 18 × 4 × 4 = 288 bin, biểu diễn tỉ lệ xuất hiện của các tổ hợp màu trong vùng đối tượng. Histogram HSV không gian sử dụng cùng cách lượng tử hóa cho bốn ô lưới, tạo 4 × 288 = 1.152 giá trị và giữ thêm thông tin phân bố màu theo vị trí. Mỗi histogram được chia cho số pixel hợp lệ trong vùng trước khi chuẩn hóa vector.")
p("LBP sử dụng 8 điểm lân cận, bán kính 1 và phương pháp uniform. Mỗi ô lưới có histogram 10 bin, tạo vector 4 × 4 × 10 = 160 chiều. HOG sử dụng 9 hướng gradient, ô 16 × 16 pixel, khối 2 × 2 ô và chuẩn hóa L2-Hys, tạo vector 7 × 7 × 2 × 2 × 9 = 1.764 chiều.")
p("Gabor gồm 16 bộ lọc được tạo từ bốn bước sóng 4, 8, 16, 32 pixel và bốn hướng 0°, 45°, 90°, 135°. Với mỗi bộ lọc, hệ thống lấy trung bình và độ lệch chuẩn của trị tuyệt đối đáp ứng trong từng ô lưới. Vector Gabor thu được có 4 × 4 × 16 × 2 = 512 chiều. Mặt nạ được co trước khi lấy thống kê LBP và Gabor để hạn chế ảnh hưởng của nền ở gần biên đối tượng; HOG được tính trên toàn ảnh xám đã đặt nền bằng 0.")
p("Cách tổ chức này kết hợp thông tin màu sắc, bố cục màu, cấu trúc biên và kết cấu bề mặt. Các nhóm mô tả cung cấp những góc nhìn khác nhau về ảnh, hỗ trợ so sánh cả hình dáng và hoa văn của động vật.")
weight_intro = p("Sau bước trích rút, hệ thống chuẩn hóa L2 độc lập cho từng nhóm trước khi ghép nối. Việc này đưa các vector hợp lệ về cùng chuẩn, hạn chế chênh lệch độ lớn giữa các nhóm và cho phép điều khiển mức đóng góp bằng trọng số. Cấu hình hiện tại ưu tiên Gabor với trọng số 0,50 để nhấn mạnh kết cấu đa hướng và đa tỉ lệ. Phân bố trọng số được trình bày trong Bảng 3.1.")
weight_intro.paragraph_format.page_break_before = True

headers = ["Nhóm\nđặc trưng", "Trọng\nsố", "Vai trò chính", "Giải thích phân bố"]
rows = [
    ("Gabor Spatial", "gabor_spatial", "Kết cấu đa tỉ lệ và đa hướng", "Được ưu tiên để mô tả sọc, đốm và vân lông qua phân bố độ lớn đáp ứng lọc theo vùng."),
    ("HOG", "hog", "Hướng gradient và cấu trúc biên", "Bổ sung thông tin về biên và cấu trúc hình dáng khi các ảnh có kết cấu gần nhau."),
    ("LBP Spatial", "lbp_spatial", "Kết cấu cục bộ", "Bổ sung mô tả các mẫu sáng tối lân cận trên bề mặt, kết hợp thông tin vị trí qua lưới 4 × 4."),
    ("HSV Spatial", "hsv_spatial", "Màu sắc theo vị trí", "Giữ phân bố màu giữa các vùng, hỗ trợ phân biệt ảnh có tổng thể màu tương tự nhưng bố cục khác nhau."),
    ("HSV Global", "hsv_global", "Thành phần màu toàn cục", "Cung cấp thông tin màu tổng thể với trọng số nhỏ, bổ trợ cho cấu trúc biên và kết cấu."),
]
widths = [3.1, 1.55, 4.1, 7.75]
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
for col, width in zip(table.columns, widths):
    col.width = Cm(width)
for cell, text in zip(table.rows[0].cells, headers):
    cell.text = text
for label, key, role, reason in rows:
    cells = table.add_row().cells
    for cell, text in zip(cells, (label, f"{DEFAULT_CONFIG.weight_map[key]:.2f}".replace(".", ","), role, reason)):
        cell.text = text
for row_index, row in enumerate(table.rows):
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))
    if row_index == 0:
        properties.append(OxmlElement("w:tblHeader"))
    for col_index, cell in enumerate(row.cells):
        cell.width = Cm(widths[col_index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_pr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            border = OxmlElement("w:" + edge)
            for attr, value in (("val", "single"), ("sz", "5"), ("color", "000000")):
                border.set(qn("w:" + attr), value)
            borders.append(border)
        cell_pr.append(borders)
        margins = OxmlElement("w:tcMar")
        for edge in ("top", "left", "bottom", "right"):
            value = OxmlElement("w:" + edge)
            value.set(qn("w:w"), "80")
            value.set(qn("w:type"), "dxa")
            margins.append(value)
        cell_pr.append(margins)
        for para in cell.paragraphs:
            para.paragraph_format.first_line_indent = Cm(0)
            para.paragraph_format.line_spacing = 1.05
            para.paragraph_format.space_after = Pt(0)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 or col_index == 1 else WD_ALIGN_PARAGRAPH.LEFT
            for r in para.runs:
                r.font.size = Pt(12)
                r.bold = row_index == 0
caption = p("Bảng 3.1: Trọng số của các nhóm đặc trưng trong cấu hình hiện tại", "Caption")
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
p("Các trọng số trên thể hiện mức ưu tiên được đặt trong cấu hình của dự án; đây chưa phải kết luận về bộ trọng số tối ưu nếu chưa có đánh giá trên tập kiểm thử độc lập.")
p("Với mỗi nhóm đặc trưng fᵢ, hệ thống tính chuẩn L2 như sau:")
eq(["‖", sub("f", "i"), "‖₂ = ", root(summation("j = 1", sub("d", "i"), sup(sub("f", "ij"), "2")))])
p("Trong đó, dᵢ là số chiều của nhóm đặc trưng thứ i và fᵢⱼ là thành phần thứ j của vector. Với vector có chuẩn lớn hơn ngưỡng ε, phép chuẩn hóa được thực hiện theo công thức:")
eq([sub(hat("f"), "i"), " = ", frac(sub("f", "i"), ["‖", sub("f", "i"), "‖₂"]), " ,     ‖", sub("f", "i"), "‖₂ > ε"])
p("Ngưỡng ε = 10⁻¹² được dùng để phát hiện vector có chuẩn quá nhỏ và tránh chia cho 0. Nếu chuẩn không vượt ngưỡng, hệ thống trả vector 0 và đánh dấu nhóm đặc trưng không hợp lệ. Vector kết hợp chỉ được tạo khi cả năm nhóm đều hợp lệ; các vector cũng được kiểm tra đủ số chiều, không âm và không chứa NaN hoặc vô cực.")
vector_intro = p("Vector cuối cùng được tạo bằng cách nhân từng vector đã chuẩn hóa với căn bậc hai trọng số rồi ghép nối theo thứ tự HSV Global, HSV Spatial, LBP Spatial, HOG và Gabor Spatial:")
vector_intro.paragraph_format.keep_with_next = True
vector_formula = eq(["v = [", root(sub("w", "1")), sub(hat("f"), "1"), " ‖ ", root(sub("w", "2")), sub(hat("f"), "2"), " ‖ ⋯ ‖ ", root(sub("w", "5")), sub(hat("f"), "5"), "]"])
vector_formula.paragraph_format.keep_with_next = True
p("Trong đó:")
bullet("f̂ᵢ là vector đã chuẩn hóa L2 của nhóm đặc trưng thứ i.")
bullet("wᵢ là trọng số của nhóm tương ứng và tổng năm trọng số bằng 1.")
bullet("Ký hiệu ‖ trong biểu thức ghép nối biểu thị việc nối các vector thành một vector dài hơn.")
p("Nhân với căn bậc hai trọng số giúp tích vô hướng giữa hai vector kết hợp bằng tổng có trọng số của các độ tương đồng cosine thành phần:")
eq([sup("v", "(a)"), " · ", sup("v", "(b)"), " = ", summation("i = 1", "5", [sub("w", "i"), " (", sup(sub(hat("f"), "i"), "(a)"), " · ", sup(sub(hat("f"), "i"), "(b)"), ")"])])
p("Nếu nhân trực tiếp vector thành phần với wᵢ thì mức đóng góp trong tích vô hướng sẽ là wᵢ². Vì vậy, công thức sử dụng căn trọng số phù hợp với cách tính điểm tổng hợp trong chế độ cosine của hệ thống. Với các vector thành phần hợp lệ có chuẩn bằng 1, vector kết hợp cũng có chuẩn xấp xỉ 1.")
p("Kết quả của bước này là một vector đặc trưng hợp nhất có 288 + 1.152 + 160 + 1.764 + 512 = 3.876 chiều. Biểu diễn này giữ thông tin của cả năm nhóm và phản ánh mức ưu tiên của cấu hình khi so khớp ảnh. Cùng quy trình trích rút, chuẩn hóa và thứ tự ghép được áp dụng cho ảnh trong kho và ảnh truy vấn để bảo đảm các vector có thể so sánh trực tiếp.")

doc.core_properties.title = "3.2.2. Thực hiện trích rút các đặc trưng đã chọn từ bộ dữ liệu"
doc.core_properties.subject = "Năm nhóm đặc trưng và vector hợp nhất 3876 chiều"
doc.core_properties.author = ""
doc.core_properties.keywords = "HSV, LBP, HOG, Gabor, L2, trích rút đặc trưng"
assert COMBINED_FEATURE_DIM == 3876
assert [FEATURE_DIMS[name] for name in FEATURE_ORDER] == [288, 1152, 160, 1764, 512]
assert abs(sum(DEFAULT_CONFIG.weight_map.values()) - 1.0) < 1e-12
assert not OUT.exists() or "--replace" in sys.argv, f"Output already exists: {OUT.name}"
doc.save(OUT)
print(f"Created: {OUT.name}; {OUT.stat().st_size} bytes; {len(doc.paragraphs)} paragraphs")
